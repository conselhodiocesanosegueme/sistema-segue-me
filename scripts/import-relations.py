import openpyxl
import json
import urllib.request
import time

import os

SECRET = os.environ.get('SUPABASE_SECRET_KEY', '')
REST_URL = f"{os.environ.get('NEXT_PUBLIC_SUPABASE_URL', 'https://vipmbcbnzybgbpygwcoe.supabase.co')}/rest/v1"

HEADERS = {
    'apikey': SECRET,
    'Authorization': f'Bearer {SECRET}',
    'Content-Type': 'application/json',
    'Prefer': 'resolution=merge-duplicates'
}

def load_maps():
    print('1. Carregando mapa de encontros do Supabase...')
    req = urllib.request.Request(f'{REST_URL}/encounters?select=id,legacy_id', headers=HEADERS)
    with urllib.request.urlopen(req) as resp:
        encounters = json.loads(resp.read().decode())
    enc_map = {e['legacy_id']: e['id'] for e in encounters if e.get('legacy_id')}
    print(f'   -> {len(enc_map)} encontros mapeados.')

    print('2. Carregando mapa de pessoas do Supabase (paginado)...')
    people_map = {}
    offset = 0
    page_size = 1000
    while True:
        req = urllib.request.Request(
            f'{REST_URL}/people?select=id,legacy_id',
            headers={**HEADERS, 'Range': f'{offset}-{offset + page_size - 1}'}
        )
        with urllib.request.urlopen(req) as resp:
            batch = json.loads(resp.read().decode())
            if not batch:
                break
            for p in batch:
                if p.get('legacy_id'):
                    people_map[p['legacy_id']] = p['id']
            offset += len(batch)
            if len(batch) < page_size:
                break
    print(f'   -> {len(people_map)} pessoas mapeadas.')
    return enc_map, people_map

def import_couples(wb, people_map):
    print('\n3. Importando Casais...')
    ws = wb['Casais']
    couples = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i < 7:
            continue
        legacy_id = str(row[0]).strip() if row[0] else None
        p1 = str(row[1]).strip() if row[1] else None
        p2 = str(row[3]).strip() if row[3] else None
        if not legacy_id or not p1 or not p2:
            continue
        if p1 not in people_map or p2 not in people_map:
            continue
        p1_id = people_map[p1]
        p2_id = people_map[p2]
        if p1_id == p2_id:
            continue
        
        start_text = str(row[5]).strip() if row[5] else None
        end_text = str(row[6]).strip() if row[6] else None
        notes = str(row[7]).strip() if row[7] else None

        couples.append({
            'legacy_id': legacy_id,
            'person_1_id': p1_id,
            'person_2_id': p2_id,
            'start_text': start_text,
            'end_text': end_text,
            'notes': notes
        })

    print(f'   -> Total de casais válidos: {len(couples)}')
    batch_size = 500
    for i in range(0, len(couples), batch_size):
        batch = couples[i:i + batch_size]
        data = json.dumps(batch).encode('utf-8')
        req = urllib.request.Request(f'{REST_URL}/couples?on_conflict=legacy_id', data=data, headers=HEADERS, method='POST')
        try:
            with urllib.request.urlopen(req) as resp:
                print(f'   Casais Lote {i // batch_size + 1}/{(len(couples) + batch_size - 1) // batch_size} inserido ({len(batch)} registros)')
        except urllib.error.HTTPError as e:
            print(f'   Erro em Casais lote {i // batch_size + 1}: {e.code} - {e.read().decode()}')
            break
        time.sleep(0.2)

def import_participations(wb, enc_map, people_map):
    print('\n4. Importando Participações...')
    ws = wb['Participações']
    participations = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i < 3:
            continue
        p_legacy = str(row[0]).strip() if row[0] else None
        e_legacy = str(row[2]).strip() if row[2] else None
        kind = str(row[3]).strip() if row[3] else 'Trabalhou'
        condition = str(row[4]).strip() if row[4] else ''
        team = str(row[5]).strip() if row[5] else None
        circle = str(row[6]).strip() if row[6] else None
        role = str(row[7]).strip() if row[7] else None
        patron = str(row[8]).strip() if row[8] else None
        source_page = str(row[9]).strip() if row[9] else None

        if not p_legacy or p_legacy not in people_map:
            continue
        if not e_legacy or e_legacy not in enc_map:
            continue

        participations.append({
            'person_id': people_map[p_legacy],
            'encounter_id': enc_map[e_legacy],
            'kind': kind,
            'condition': condition,
            'team': team,
            'circle': circle,
            'role': role,
            'patron': patron,
            'source_page': source_page
        })

    print(f'   -> Total de participações válidas: {len(participations)}')
    batch_size = 1000
    for i in range(0, len(participations), batch_size):
        batch = participations[i:i + batch_size]
        data = json.dumps(batch).encode('utf-8')
        req = urllib.request.Request(f'{REST_URL}/participations', data=data, headers=HEADERS, method='POST')
        try:
            with urllib.request.urlopen(req) as resp:
                if (i // batch_size + 1) % 5 == 0 or (i + batch_size >= len(participations)):
                    print(f'   Participações Lote {i // batch_size + 1}/{(len(participations) + batch_size - 1) // batch_size} inserido ({len(batch)} registros)')
        except urllib.error.HTTPError as e:
            print(f'   Erro em Participações lote {i // batch_size + 1}: {e.code} - {e.read().decode()}')
            break
        time.sleep(0.2)

def import_mandates(wb, enc_map, people_map):
    print('\n5. Importando Mandatos...')
    ws = wb['Mandatos']
    mandates = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i < 3:
            continue
        p_legacy = str(row[0]).strip() if row[0] else None
        e_legacy = str(row[2]).strip() if row[2] else None
        body = str(row[3]).strip() if row[3] else None
        role = str(row[4]).strip() if row[4] else 'Membro'
        condition = str(row[5]).strip() if row[5] else None
        s_year_val = row[6]
        e_year_val = row[7]
        rec_type = str(row[8]).strip() if row[8] else None
        notes = str(row[9]).strip() if row[9] else None

        if not p_legacy or p_legacy not in people_map:
            continue
        if not body:
            continue

        enc_id = enc_map.get(e_legacy) if e_legacy else None

        s_year = int(s_year_val) if s_year_val and str(s_year_val).strip().isdigit() else None
        e_year = int(e_year_val) if e_year_val and str(e_year_val).strip().isdigit() else None
        if s_year and e_year and e_year < s_year:
            e_year = None

        mandates.append({
            'person_id': people_map[p_legacy],
            'encounter_id': enc_id,
            'body': body,
            'role': role,
            'condition': condition,
            'start_year': s_year,
            'end_year': e_year,
            'record_type': rec_type,
            'notes': notes
        })

    print(f'   -> Total de mandatos válidos: {len(mandates)}')
    batch_size = 1000
    for i in range(0, len(mandates), batch_size):
        batch = mandates[i:i + batch_size]
        data = json.dumps(batch).encode('utf-8')
        req = urllib.request.Request(f'{REST_URL}/mandates', data=data, headers=HEADERS, method='POST')
        try:
            with urllib.request.urlopen(req) as resp:
                print(f'   Mandatos Lote {i // batch_size + 1}/{(len(mandates) + batch_size - 1) // batch_size} inserido ({len(batch)} registros)')
        except urllib.error.HTTPError as e:
            print(f'   Erro em Mandatos lote {i // batch_size + 1}: {e.code} - {e.read().decode()}')
            break
        time.sleep(0.2)

def main():
    wb = openpyxl.load_workbook('BACKUP - Base Historica Segue-me - apos 2023-SFA-SFG-12 - 20260911-1546.xlsx', read_only=True)
    enc_map, people_map = load_maps()
    import_couples(wb, people_map)
    import_participations(wb, enc_map, people_map)
    import_mandates(wb, enc_map, people_map)
    print('\n=== MIGRAÇÃO DAS RELAÇÕES CONCLUÍDA COM SUCESSO! ===')

if __name__ == '__main__':
    main()
