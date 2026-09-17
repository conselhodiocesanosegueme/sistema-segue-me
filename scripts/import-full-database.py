#!/usr/bin/env python3
"""
Sistema Segue-me - Script Unificado de Importação da Planilha Oficial
===================================================================
Este script processa a planilha oficial (.xlsx), valida a integridade de todos os IDs
e realiza o upsert/sincronização idempotente nas tabelas do Supabase:
  1. Encontros (upsert por legacy_id)
  2. Pessoas (upsert por legacy_id)
  3. Casais (upsert por legacy_id)
  4. Participações (deduplicação inteligente - nunca duplica)
  5. Mandatos (deduplicação inteligente - nunca duplica)

Uso:
  python3 scripts/import-full-database.py --dry-run          # Simula e valida IDs sem gravar no banco
  python3 scripts/import-full-database.py                    # Executa a importação completa
  python3 scripts/import-full-database.py --file caminho.xlsx # Especifica um arquivo alternativo
"""

import sys
import os
import glob
import json
import time
import argparse
import urllib.request
import urllib.error
import openpyxl

# Carrega variáveis de ambiente do .env.local se existirem
def load_env():
    env_path = os.path.join(os.path.dirname(__file__), '..', '.env.local')
    if os.path.exists(env_path):
        with open(env_path, 'r', encoding='utf-8') as f:
            for line in f:
                line = line.strip()
                if line and not line.startswith('#') and '=' in line:
                    k, v = line.split('=', 1)
                    if k not in os.environ:
                        os.environ[k] = v

load_env()

SECRET = os.environ.get('SUPABASE_SECRET_KEY', '')
SUPABASE_URL = os.environ.get('NEXT_PUBLIC_SUPABASE_URL', 'https://vipmbcbnzybgbpygwcoe.supabase.co')
REST_URL = f"{SUPABASE_URL}/rest/v1"

HEADERS = {
    'apikey': SECRET,
    'Authorization': f'Bearer {SECRET}',
    'Content-Type': 'application/json',
}

def find_default_xlsx():
    root = os.path.join(os.path.dirname(__file__), '..')
    files = sorted(glob.glob(os.path.join(root, '*.xlsx')), key=os.path.getmtime, reverse=True)
    if files:
        return files[0]
    return None

def api_get(endpoint, params=''):
    url = f"{REST_URL}/{endpoint}?{params}" if params else f"{REST_URL}/{endpoint}"
    req = urllib.request.Request(url, headers=HEADERS)
    with urllib.request.urlopen(req) as resp:
        return json.loads(resp.read().decode())

def api_post(endpoint, data, upsert=False):
    url = f"{REST_URL}/{endpoint}"
    headers = {**HEADERS}
    if upsert:
        headers['Prefer'] = 'resolution=merge-duplicates'
    
    payload = json.dumps(data).encode('utf-8')
    req = urllib.request.Request(url, data=payload, headers=headers, method='POST')
    with urllib.request.urlopen(req) as resp:
        return resp.status

def load_db_maps():
    print('📦 Carregando mapas de entidades existentes no Supabase...')
    
    # 1. Encontros
    enc_data = api_get('encounters', 'select=id,legacy_id')
    enc_map = {e['legacy_id']: e['id'] for e in enc_data if e.get('legacy_id')}
    print(f'   • {len(enc_map)} encontros mapeados no banco.')

    # 2. Pessoas (paginado)
    people_map = {}
    offset = 0
    page_size = 1000
    while True:
        req = urllib.request.Request(
            f'{REST_URL}/people?select=id,legacy_id',
            headers={**HEADERS, 'Range': f'{offset}-{offset + page_size - 1}'}
        )
        with urllib.request.urlopen(req) as resp:
            batch = json.loads(resp.read().decode())
            if not batch:
                break
            for p in batch:
                if p.get('legacy_id'):
                    people_map[p['legacy_id']] = p['id']
            offset += len(batch)
            if len(batch) < page_size:
                break
    print(f'   • {len(people_map)} pessoas mapeadas no banco.')

    return enc_map, people_map

def import_encounters(wb, dry_run=False):
    print('\n🏛️ [1/5] Processando Encontros...')
    if 'Encontros' not in wb.sheetnames:
        print('   ⚠️ Aba "Encontros" não encontrada.')
        return {}

    ws = wb['Encontros']
    rows = list(ws.iter_rows(values_only=True))
    encounters = []
    
    # Linha 3 (índice 2) contém cabeçalho, dados começam na linha 4 (índice 3)
    for r in rows[3:]:
        if not r or not r[0]:
            continue
        legacy_id = str(r[0]).strip()
        edition = str(r[1]).strip() if r[1] is not None else ''
        year_val = r[2]
        year = int(year_val) if year_val and str(year_val).strip().isdigit() else None
        parish = str(r[3]).strip() if r[3] is not None else ''
        city = str(r[4]).strip() if r[4] is not None else ''
        name = str(r[5]).strip() if r[5] is not None else ''
        date_text = str(r[6]).strip() if r[6] is not None else ''
        source_name = str(r[7]).strip() if len(r) > 7 and r[7] is not None else None
        extraction_status = str(r[8]).strip() if len(r) > 8 and r[8] is not None else 'Consolidado'
        notes = str(r[9]).strip() if len(r) > 9 and r[9] is not None else None
        target_diocese = str(r[10]).strip() if len(r) > 10 and r[10] is not None else None

        # Detecção de Implantação Externa (Outra Diocese apadrinhada por Anápolis)
        is_external = False
        combined_text = f"{name} {parish} {city} {notes or ''}".lower()
        if 'implantação externa' in combined_text or 'implantacao externa' in combined_text or 'outra diocese' in combined_text:
            is_external = True
        
        external_dioceses = [
            'Uruaçu', 'Rubiataba', 'Mozarlândia', 'Luziânia', 'Goiânia', 'Palmas',
            'Formosa', 'Tocantinópolis', 'Ipameri', 'Itumbiara', 'Jataí', 'Rio Verde'
        ]
        for d in external_dioceses:
            if d.lower() in combined_text:
                is_external = True
                if not target_diocese:
                    target_diocese = f"Diocese de {d}" if not d in ['Goiânia', 'Palmas'] else f"Arquidiocese de {d}"
                break

        enc_record = {
            'legacy_id': legacy_id,
            'edition': edition,
            'year': year,
            'parish': parish,
            'city': city,
            'name': name,
            'date_text': date_text,
            'source_name': source_name,
            'extraction_status': extraction_status,
            'notes': notes,
            'is_external_implantation': is_external,
            'target_diocese': target_diocese,
        }

        encounters.append(enc_record)

    print(f'   • Total de encontros na planilha: {len(encounters)}')
    if dry_run:
        print('   🔍 [DRY-RUN] Encontros validados com sucesso (nenhuma gravação realizada).')
        return encounters

    batch_size = 100
    for i in range(0, len(encounters), batch_size):
        batch = encounters[i:i + batch_size]
        url = f"{REST_URL}/encounters?on_conflict=legacy_id"
        headers = {**HEADERS, 'Prefer': 'resolution=merge-duplicates'}
        req = urllib.request.Request(url, data=json.dumps(batch).encode('utf-8'), headers=headers, method='POST')
        with urllib.request.urlopen(req) as resp:
            pass
    print('   ✅ Encontros gravados/atualizados com sucesso.')
    return encounters

def import_people(wb, dry_run=False):
    print('\n👤 [2/5] Processando Pessoas...')
    if 'Pessoas' not in wb.sheetnames:
        print('   ⚠️ Aba "Pessoas" não encontrada.')
        return []

    ws = wb['Pessoas']
    rows = list(ws.iter_rows(values_only=True))
    people = []

    for r in rows[3:]:
        if not r or not r[1]: # Precisa de nome
            continue
        legacy_id = str(r[0]).strip() if r[0] else None
        name = str(r[1]).strip()
        phone = str(r[2]).strip() if r[2] else None
        email = str(r[3]).strip() if r[3] else None
        birth_date_text = str(r[4]).strip() if r[4] else None
        sex = str(r[5]).strip() if r[5] else None
        identification_status = str(r[8]).strip() if len(r) > 8 and r[8] else 'Identificado'
        notes = str(r[9]).strip() if len(r) > 9 and r[9] else None

        people.append({
            'legacy_id': legacy_id,
            'name': name,
            'phone': phone,
            'email': email,
            'birth_date_text': birth_date_text,
            'sex': sex,
            'identification_status': identification_status,
            'notes': notes,
        })

    print(f'   • Total de pessoas na planilha: {len(people)}')
    if dry_run:
        print('   🔍 [DRY-RUN] Pessoas validadas com sucesso (nenhuma gravação realizada).')
        return people

    batch_size = 1000
    total = len(people)
    for i in range(0, total, batch_size):
        batch = people[i:i + batch_size]
        url = f"{REST_URL}/people?on_conflict=legacy_id"
        headers = {**HEADERS, 'Prefer': 'resolution=merge-duplicates'}
        req = urllib.request.Request(url, data=json.dumps(batch).encode('utf-8'), headers=headers, method='POST')
        with urllib.request.urlopen(req) as resp:
            pass
        if (i // batch_size + 1) % 5 == 0 or i + batch_size >= total:
            print(f'   Lote {i // batch_size + 1}/{(total + batch_size - 1) // batch_size} inserido/atualizado ({len(batch)} registros)')
        time.sleep(0.1)

    print('   ✅ Pessoas gravadas/atualizados com sucesso.')
    return people

def import_couples(wb, people_map, dry_run=False):
    print('\n💍 [3/5] Processando Casais...')
    if 'Casais' not in wb.sheetnames:
        print('   ⚠️ Aba "Casais" não encontrada.')
        return []

    ws = wb['Casais']
    couples = []
    missing_people = set()

    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i < 3 or not row or not row[0]:
            continue
        legacy_id = str(row[0]).strip()
        p1 = str(row[1]).strip() if row[1] else None
        p2 = str(row[3]).strip() if len(row) > 3 and row[3] else None

        if not legacy_id or not p1 or not p2:
            continue
        
        p1_id = people_map.get(p1)
        p2_id = people_map.get(p2)

        if not p1_id:
            missing_people.add(p1)
        if not p2_id:
            missing_people.add(p2)

        if not p1_id or not p2_id or p1_id == p2_id:
            continue

        start_text = str(row[5]).strip() if len(row) > 5 and row[5] else None
        end_text = str(row[6]).strip() if len(row) > 6 and row[6] else None
        notes = str(row[7]).strip() if len(row) > 7 and row[7] else None

        couples.append({
            'legacy_id': legacy_id,
            'person_1_id': p1_id,
            'person_2_id': p2_id,
            'start_text': start_text,
            'end_text': end_text,
            'notes': notes,
        })

    print(f'   • Total de casais válidos: {len(couples)}')
    if missing_people:
        print(f'   ⚠️ {len(missing_people)} IDs de pessoas citadas em Casais não foram localizadas em Pessoas.')

    if dry_run:
        print('   🔍 [DRY-RUN] Casais validados com sucesso.')
        return couples

    batch_size = 500
    for i in range(0, len(couples), batch_size):
        batch = couples[i:i + batch_size]
        url = f"{REST_URL}/couples?on_conflict=legacy_id"
        headers = {**HEADERS, 'Prefer': 'resolution=merge-duplicates'}
        req = urllib.request.Request(url, data=json.dumps(batch).encode('utf-8'), headers=headers, method='POST')
        with urllib.request.urlopen(req) as resp:
            pass
        time.sleep(0.1)

    print('   ✅ Casais gravados/atualizados com sucesso.')
    return couples

def fetch_existing_participations_signatures():
    print('   📥 Carregando participações já existentes para evitar duplicações...')
    signatures = set()
    offset = 0
    page_size = 2000
    while True:
        req = urllib.request.Request(
            f'{REST_URL}/participations?select=person_id,encounter_id,kind,team,circle,role',
            headers={**HEADERS, 'Range': f'{offset}-{offset + page_size - 1}'}
        )
        try:
            with urllib.request.urlopen(req) as resp:
                batch = json.loads(resp.read().decode())
                if not batch:
                    break
                for row in batch:
                    sig = (
                        row['person_id'],
                        row['encounter_id'],
                        row['kind'],
                        row.get('team') or '',
                        row.get('circle') or '',
                        row.get('role') or '',
                    )
                    signatures.add(sig)
                offset += len(batch)
                if len(batch) < page_size:
                    break
        except Exception as e:
            print(f'   ⚠️ Falha ao ler assinaturas de participações: {e}')
            break
    print(f'   • {len(signatures)} participações existentes indexadas.')
    return signatures

def import_participations(wb, enc_map, people_map, dry_run=False):
    print('\n🤝 [4/5] Processando Participações...')
    if 'Participações' not in wb.sheetnames:
        print('   ⚠️ Aba "Participações" não encontrada.')
        return []

    ws = wb['Participações']
    participations_to_insert = []
    missing_p = set()
    missing_e = set()

    existing_sigs = set() if dry_run else fetch_existing_participations_signatures()

    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i < 3 or not row or not row[0]:
            continue
        p_legacy = str(row[0]).strip()
        e_legacy = str(row[2]).strip() if len(row) > 2 and row[2] else None
        kind = str(row[3]).strip() if len(row) > 3 and row[3] else 'Trabalhou'
        condition = str(row[4]).strip() if len(row) > 4 and row[4] else ''
        team = str(row[5]).strip() if len(row) > 5 and row[5] else None
        circle = str(row[6]).strip() if len(row) > 6 and row[6] else None
        role = str(row[7]).strip() if len(row) > 7 and row[7] else None
        patron = str(row[8]).strip() if len(row) > 8 and row[8] else None
        source_page = str(row[9]).strip() if len(row) > 9 and row[9] else None
        notes_part = str(row[10]).strip() if len(row) > 10 and row[10] else None

        is_seed = False
        if notes_part and ('remessa' in notes_part.lower() or 'fora' in notes_part.lower() or 'implanta' in notes_part.lower()):
            is_seed = True

        p_id = people_map.get(p_legacy)
        e_id = enc_map.get(e_legacy)

        if not p_id:
            missing_p.add(p_legacy)
            continue
        if not e_id:
            missing_e.add(e_legacy)
            continue

        sig = (p_id, e_id, kind, team or '', circle or '', role or '')
        if sig in existing_sigs:
            continue # Já gravado exatamente igual, evita duplicidade!

        participations_to_insert.append({
            'person_id': p_id,
            'encounter_id': e_id,
            'kind': kind,
            'condition': condition,
            'team': team,
            'circle': circle,
            'role': role,
            'patron': patron,
            'source_page': source_page,
            'is_external_seed': is_seed,
        })
        existing_sigs.add(sig)

    print(f'   • Novas participações a inserir: {len(participations_to_insert)}')
    if missing_p:
        print(f'   ⚠️ {len(missing_p)} pessoas referenciadas em Participações não foram encontradas em Pessoas.')
    if missing_e:
        print(f'   ⚠️ {len(missing_e)} encontros referenciados em Participações não foram encontrados em Encontros.')

    if dry_run:
        print('   🔍 [DRY-RUN] Participações validadas com sucesso.')
        return participations_to_insert

    batch_size = 1000
    total = len(participations_to_insert)
    for i in range(0, total, batch_size):
        batch = participations_to_insert[i:i + batch_size]
        url = f"{REST_URL}/participations"
        req = urllib.request.Request(url, data=json.dumps(batch).encode('utf-8'), headers=HEADERS, method='POST')
        with urllib.request.urlopen(req) as resp:
            pass
        if (i // batch_size + 1) % 5 == 0 or i + batch_size >= total:
            print(f'   Lote {i // batch_size + 1}/{(total + batch_size - 1) // batch_size} inserido ({len(batch)} registros)')
        time.sleep(0.1)

    print('   ✅ Novas participações gravadas com sucesso.')
    return participations_to_insert

def fetch_existing_mandates_signatures():
    print('   📥 Carregando mandatos existentes para evitar duplicações...')
    signatures = set()
    offset = 0
    page_size = 1000
    while True:
        req = urllib.request.Request(
            f'{REST_URL}/mandates?select=person_id,body,role,start_year',
            headers={**HEADERS, 'Range': f'{offset}-{offset + page_size - 1}'}
        )
        try:
            with urllib.request.urlopen(req) as resp:
                batch = json.loads(resp.read().decode())
                if not batch:
                    break
                for row in batch:
                    sig = (
                        row['person_id'],
                        row.get('body') or '',
                        row.get('role') or '',
                        row.get('start_year'),
                    )
                    signatures.add(sig)
                offset += len(batch)
                if len(batch) < page_size:
                    break
        except Exception as e:
            print(f'   ⚠️ Falha ao ler assinaturas de mandatos: {e}')
            break
    print(f'   • {len(signatures)} mandatos existentes indexados.')
    return signatures

def import_mandates(wb, enc_map, people_map, dry_run=False):
    print('\n📜 [5/5] Processando Mandatos...')
    if 'Mandatos' not in wb.sheetnames:
        print('   ⚠️ Aba "Mandatos" não encontrada.')
        return []

    ws = wb['Mandatos']
    mandates_to_insert = []
    missing_p = set()

    existing_sigs = set() if dry_run else fetch_existing_mandates_signatures()

    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i < 3 or not row or not row[0]:
            continue
        p_legacy = str(row[0]).strip()
        e_legacy = str(row[2]).strip() if len(row) > 2 and row[2] else None
        body = str(row[3]).strip() if len(row) > 3 and row[3] else None
        role = str(row[4]).strip() if len(row) > 4 and row[4] else 'Membro'
        condition = str(row[5]).strip() if len(row) > 5 and row[5] else None
        s_year_val = row[6] if len(row) > 6 else None
        e_year_val = row[7] if len(row) > 7 else None
        rec_type = str(row[8]).strip() if len(row) > 8 and row[8] else None
        notes = str(row[9]).strip() if len(row) > 9 and row[9] else None

        p_id = people_map.get(p_legacy)
        if not p_id:
            missing_p.add(p_legacy)
            continue
        if not body:
            continue

        enc_id = enc_map.get(e_legacy) if e_legacy else None

        s_year = int(s_year_val) if s_year_val and str(s_year_val).strip().isdigit() else None
        e_year = int(e_year_val) if e_year_val and str(e_year_val).strip().isdigit() else None
        if s_year and e_year and e_year < s_year:
            e_year = None

        sig = (p_id, body, role, s_year)
        if sig in existing_sigs:
            continue # Evita duplicação

        mandates_to_insert.append({
            'person_id': p_id,
            'encounter_id': enc_id,
            'body': body,
            'role': role,
            'condition': condition,
            'start_year': s_year,
            'end_year': e_year,
            'record_type': rec_type,
            'notes': notes,
        })
        existing_sigs.add(sig)

    print(f'   • Novos mandatos a inserir: {len(mandates_to_insert)}')
    if missing_p:
        print(f'   ⚠️ {len(missing_p)} pessoas referenciadas em Mandatos não foram encontradas em Pessoas.')

    if dry_run:
        print('   🔍 [DRY-RUN] Mandatos validados com sucesso.')
        return mandates_to_insert

    batch_size = 1000
    total = len(mandates_to_insert)
    for i in range(0, total, batch_size):
        batch = mandates_to_insert[i:i + batch_size]
        url = f"{REST_URL}/mandates"
        req = urllib.request.Request(url, data=json.dumps(batch).encode('utf-8'), headers=HEADERS, method='POST')
        with urllib.request.urlopen(req) as resp:
            pass
        print(f'   Lote {i // batch_size + 1}/{(total + batch_size - 1) // batch_size} inserido ({len(batch)} registros)')
        time.sleep(0.1)

    print('   ✅ Novos mandatos gravados com sucesso.')
    return mandates_to_insert

def main():
    parser = argparse.ArgumentParser(description='Importador Unificado do Sistema Segue-me')
    parser.add_argument('--file', help='Caminho do arquivo .xlsx da planilha')
    parser.add_argument('--dry-run', action='store_true', help='Apenas simula a leitura e valida integridade de IDs sem gravar')
    args = parser.parse_args()

    xlsx_path = args.file or find_default_xlsx()
    if not xlsx_path or not os.path.exists(xlsx_path):
        print(f'❌ Arquivo Excel não encontrado: {xlsx_path}')
        sys.exit(1)

    print('=' * 70)
    print('  SISTEMA SEGUE-ME - SINCRONIZADOR OFICIAL DE BASE HISTÓRICA')
    print('=' * 70)
    print(f'📄 Arquivo: {os.path.basename(xlsx_path)}')
    print(f'⚙️  Modo: {"🔍 SIMULAÇÃO (DRY-RUN - sem gravação)" if args.dry_run else "🚀 PRODUÇÃO (atualização do banco)"}')
    print('=' * 70)

    start_time = time.time()
    wb = openpyxl.load_workbook(xlsx_path, read_only=True)

    # 1. Encontros
    import_encounters(wb, dry_run=args.dry_run)

    # 2. Pessoas
    import_people(wb, dry_run=args.dry_run)

    # Carrega/atualiza os mapas com IDs atuais no Supabase
    enc_map, people_map = load_db_maps()

    # 3. Casais
    import_couples(wb, people_map, dry_run=args.dry_run)

    # 4. Participações
    import_participations(wb, enc_map, people_map, dry_run=args.dry_run)

    # 5. Mandatos
    import_mandates(wb, enc_map, people_map, dry_run=args.dry_run)

    elapsed = time.time() - start_time
    print('\n' + '=' * 70)
    print(f'🎉 PROCESSO CONCLUÍDO COM SUCESSO EM {elapsed:.1f}s!')
    print('=' * 70)

if __name__ == '__main__':
    main()
